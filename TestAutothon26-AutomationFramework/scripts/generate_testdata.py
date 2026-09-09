"""Creates testdata/test_data.xlsx - run once after `pip install -r requirements.txt`.

The workbook is the single source of test data required by the challenge.
Real mobile numbers / OTPs stay in .env, never in this file.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook

from config.settings import TEST_DATA_DIR

SHEETS: dict[str, tuple[list[str], list[list]]] = {
    "Journey": (
        ["TestCaseId", "Pincode", "Category", "Brand", "SecondFilter", "ProductName",
         "BargainAttempts", "PaymentMethod", "AddressName", "AddressPhone",
         "AddressLine", "City"],
        [[
            "TC_E2E_01", "560025", "Toys & Games", "Kidsnation", "",
            "Classic 15.7 Inch Soft Tip Dartboard Game Set", 3, "Net Banking",
            "Test Autothon", "9999999999", "901 KP Aurum, Marol", "Mumbai",
        ]],
    ),
    "Login": (
        ["TestCaseId", "Scenario", "MobileNumber", "Otp", "ExpectedResult"],
        [
            ["TC_LOGIN_01", "valid_mobile_and_default_otp", "9999999999", "123456", "Pass"],
            ["TC_LOGIN_02", "invalid_otp", "9999999999", "111111", "Fail"],
            ["TC_LOGIN_03", "mobile_too_short", "12345", "", "Fail"],
            ["TC_LOGIN_04", "alphabetic_mobile", "abcdefghij", "", "Fail"],
            ["TC_LOGIN_05", "empty_mobile", "", "", "Fail"],
        ],
    ),
    "Pincode": (
        ["TestCaseId", "Scenario", "Pincode", "ExpectedResult"],
        [
            ["TC_PIN_01", "serviceable_bengaluru", "560025", "Pass"],
            ["TC_PIN_02", "serviceable_mumbai", "400059", "Pass"],
            ["TC_PIN_03", "invalid_five_digits", "56002", "Fail"],
            ["TC_PIN_04", "non_numeric", "ABC123", "Fail"],
        ],
    ),
}


def main() -> None:
    TEST_DATA_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for name, (headers, rows) in SHEETS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for idx, header in enumerate(headers, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=idx).column_letter].width = (
                max(len(header) + 4, 18)
            )

    target = TEST_DATA_DIR / "test_data.xlsx"
    workbook.save(target)
    print(f"Created {target}")


if __name__ == "__main__":
    main()
