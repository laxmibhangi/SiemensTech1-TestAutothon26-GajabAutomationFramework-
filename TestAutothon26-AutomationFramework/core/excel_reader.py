"""Excel-backed test data provider.

The challenge mandates that test data is sourced from an Excel file. Each sheet
is read into a list of dicts keyed by the header row, which plugs straight into
``pytest.mark.parametrize`` for data-driven execution.
"""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from config.settings import TEST_DATA_DIR

DEFAULT_WORKBOOK = TEST_DATA_DIR / "test_data.xlsx"


@lru_cache(maxsize=None)
def _load(workbook: Path) -> dict[str, tuple[dict[str, Any], ...]]:
    if not workbook.exists():
        raise FileNotFoundError(
            f"Test data workbook not found: {workbook}. "
            "Run `python scripts/generate_testdata.py` to create it."
        )

    wb = load_workbook(workbook, data_only=True)
    data: dict[str, tuple[dict[str, Any], ...]] = {}

    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            data[sheet.title] = ()
            continue

        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        records = []
        for row in rows[1:]:
            if all(cell is None for cell in row):
                continue
            records.append({h: v for h, v in zip(headers, row) if h})
        data[sheet.title] = tuple(records)

    wb.close()
    return data


def read_sheet(sheet_name: str, workbook: Path = DEFAULT_WORKBOOK) -> list[dict[str, Any]]:
    """Return every row of ``sheet_name`` as a list of dicts."""
    sheets = _load(workbook)
    if sheet_name not in sheets:
        raise KeyError(f"Sheet '{sheet_name}' not in {workbook}. Found: {list(sheets)}")
    return [dict(r) for r in sheets[sheet_name]]


def read_row(sheet_name: str, key_column: str, key_value: Any,
             workbook: Path = DEFAULT_WORKBOOK) -> dict[str, Any]:
    """Return the first row where ``key_column`` equals ``key_value``."""
    for row in read_sheet(sheet_name, workbook):
        if str(row.get(key_column)) == str(key_value):
            return row
    raise KeyError(f"No row with {key_column}={key_value!r} in sheet '{sheet_name}'")
